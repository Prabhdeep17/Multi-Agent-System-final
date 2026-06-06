import os
import time
from pathlib import Path
import pandas as pd
from fpdf import FPDF
from docx import Document
from multimodal_processor import process_document
from vector_store import build_faiss_index
from langchain_core.documents import Document as LC_Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from agents_graph import build_graph, run_query

SESSION_ID = "comprehensive_test_888"
SESSION_DIR = Path("./session_docs") / SESSION_ID
SESSION_DIR.mkdir(parents=True, exist_ok=True)

def generate_test_files():
    print("--- Generating Test Files ---")
    
    # 1. Text File
    txt_path = SESSION_DIR / "simple.txt"
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write("This is a simple text file. The secret code is ALFA-77.")
    
    # 2. DOCX File
    docx_path = SESSION_DIR / "report.docx"
    doc = Document()
    doc.add_heading('Test Document', 0)
    doc.add_paragraph('The revenue for Q1 was $50,000.')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Product"
    table.cell(0, 1).text = "Sales"
    table.cell(1, 0).text = "Widget"
    table.cell(1, 1).text = "100"
    doc.save(docx_path)
    
    return [txt_path, docx_path]

def test_pipeline():
    api_key = os.getenv("GOOGLE_API_KEY")
    if not api_key:
        print("FAIL: No API key found.")
        return

    files = generate_test_files()
    
    print("\n--- Processing Files ---")
    all_text = ""
    for f in files:
        text = process_document(str(f), SESSION_ID, api_key)
        all_text += f"\n\nSource: {f.name}\n{text}"
    
    # Build FAISS
    splitter = RecursiveCharacterTextSplitter(chunk_size=1000, chunk_overlap=200)
    doc = LC_Document(page_content=all_text, metadata={"source": "test_suite"})
    chunks = splitter.split_documents([doc])
    faiss_dir = Path("./faiss_index") / SESSION_ID
    vs = build_faiss_index(chunks, str(faiss_dir))
    
    app = build_graph(vs, SESSION_ID)
    
    queries = [
        "Hi there, how are you today?", # Should route to General/Conversational
        "What is the secret code?", # Should route to Retrieval
        "What is the revenue for Q1 and what were the total sales for Widgets based on the table?", # Hybrid Data + Retrieval
    ]

    for q in queries:
        print(f"\n======================================")
        print(f"QUERY: {q}")
        try:
            for event in run_query(app, q, api_key=api_key, session_id=SESSION_ID):
                if "status" in event:
                    print(f"STATUS: {event['status']}")
                elif "final_answer" in event:
                    print("FINAL ANSWER:")
                    print(event["final_answer"])
                    print("\nPATH: ", " -> ".join(event.get("agent_chain", [])))
                elif "error" in event:
                    print(f"PIPELINE ERROR CAUGHT: {event.get('final_answer')}")
        except Exception as e:
            print(f"CRITICAL CRASH: {e}")

if __name__ == "__main__":
    test_pipeline()
